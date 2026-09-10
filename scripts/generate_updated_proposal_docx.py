"""
generate_updated_proposal_docx.py
==================================
Converts Proposal_Mahdi_Abdollahzadeh_Updated.md into an official, publication-quality
University of Tehran Bachelor's Thesis Proposal Word document (.docx).
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

import pypandoc

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INPUT_MD = os.path.join(PROJECT_ROOT, "docs", "thesis", "Proposal_Mahdi_Abdollahzadeh_Updated.md")
TEMP_DOCX = os.path.join(PROJECT_ROOT, "docs", "thesis", "temp_proposal.docx")
FINAL_DOCX = os.path.join(PROJECT_ROOT, "docs", "thesis", "Proposal_Mahdi_Abdollahzadeh_Updated.docx")

PERSIAN_PATTERN = re.compile(r'[\u0600-\u06FF\uFB50-\uFDFF\uFE70-\uFEFF]')

def is_persian_text(text):
    if not text:
        return False
    return len(PERSIAN_PATTERN.findall(text)) > 0

def set_p_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p.alignment = align
    pPr = p._p.get_or_add_pPr()
    for existing_bidi in pPr.findall(qn('w:bidi')):
        pPr.remove(existing_bidi)
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)

def set_p_ltr(p, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.alignment = align
    pPr = p._p.get_or_add_pPr()
    for existing_bidi in pPr.findall(qn('w:bidi')):
        pPr.remove(existing_bidi)

def set_run_font(run, persian_font="B Nazanin", latin_font="Times New Roman", size_pt=12.5, bold=None, italic=None, color_rgb=None):
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), latin_font)
    rFonts.set(qn('w:hAnsi'), latin_font)
    rFonts.set(qn('w:cs'), persian_font)

    if is_persian_text(run.text):
        for existing_rtl in rPr.findall(qn('w:rtl')):
            rPr.remove(existing_rtl)
        rtl = OxmlElement('w:rtl')
        rtl.set(qn('w:val'), '1')
        rPr.append(rtl)

def set_table_rtl_and_borders(table):
    tblPr = table._tbl.tblPr
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for bidi in tblPr.findall(qn('w:bidiVisual')):
        tblPr.remove(bidi)
    bidiVisual = OxmlElement('w:bidiVisual')
    bidiVisual.set(qn('w:val'), '1')
    tblPr.append(bidiVisual)

    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="888888"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>\n'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="888888"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_proposal_docx():
    print("Converting Markdown to base DOCX via Pandoc...")
    cur_dir = os.getcwd()
    os.chdir(PROJECT_ROOT)
    try:
        pypandoc.convert_file(
            INPUT_MD,
            'docx',
            outputfile=TEMP_DOCX,
            extra_args=['--resource-path=.']
        )
    finally:
        os.chdir(cur_dir)

    print("Styling University of Tehran Proposal DOCX...")
    doc = docx.Document(TEMP_DOCX)

    # Standard Proposal Margins: Top 2.5cm, Bottom 2.5cm, Right 2.8cm, Left 2.5cm
    for section in doc.sections:
        section.top_margin = Inches(0.98)
        section.bottom_margin = Inches(0.98)
        section.right_margin = Inches(1.10)
        section.left_margin = Inches(0.98)

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        style_name = p.style.name.lower()
        is_persian = is_persian_text(text)

        # Main Header / Title of Form
        if text.startswith("فرم معرفی پروژه کارشناسی") or text.startswith("بسمه تعالی"):
            set_p_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                set_run_font(run, persian_font="B Titr", latin_font="Times New Roman", size_pt=15.0, bold=True, color_rgb=RGBColor(20, 35, 75))
            continue

        if text.startswith("پردیس دانشکده‌های فنی"):
            set_p_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(12)
            for run in p.runs:
                set_run_font(run, persian_font="B Titr", latin_font="Times New Roman", size_pt=13.0, bold=True, color_rgb=RGBColor(50, 50, 50))
            continue

        # Headings (۱- عنوان پروژه, ۲- مشخصات, etc.)
        if 'heading 1' in style_name or 'heading 2' in style_name or 'heading 3' in style_name:
            set_p_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                set_run_font(run, persian_font="B Titr", latin_font="Times New Roman", size_pt=13.5, bold=True, color_rgb=RGBColor(25, 45, 90))

        elif 'heading 4' in style_name:
            set_p_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            for run in p.runs:
                set_run_font(run, persian_font="B Titr", latin_font="Times New Roman", size_pt=12.5, bold=True, color_rgb=RGBColor(40, 70, 120))

        else:
            # Regular paragraphs
            if is_persian:
                set_p_rtl(p, WD_ALIGN_PARAGRAPH.JUSTIFY)
                p.paragraph_format.line_spacing = 1.18
                p.paragraph_format.space_after = Pt(4)
                for run in p.runs:
                    set_run_font(run, persian_font="B Nazanin", latin_font="Times New Roman", size_pt=12.5)
            else:
                set_p_ltr(p, WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(4)
                for run in p.runs:
                    set_run_font(run, persian_font="B Nazanin", latin_font="Times New Roman", size_pt=11.5)

    # Style all tables
    for table in doc.tables:
        set_table_rtl_and_borders(table)
        for row_idx, row in enumerate(table.rows):
            is_header = (row_idx == 0)
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if is_header:
                    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EBF0F5"/>')
                    cell._tc.get_or_add_tcPr().append(shading)
                for p in cell.paragraphs:
                    p_text = p.text.strip()
                    if is_persian_text(p_text):
                        set_p_rtl(p, WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.RIGHT)
                        for run in p.runs:
                            set_run_font(run, persian_font="B Nazanin", latin_font="Times New Roman", size_pt=11.0, bold=is_header)
                    else:
                        set_p_ltr(p, WD_ALIGN_PARAGRAPH.CENTER)
                        for run in p.runs:
                            set_run_font(run, persian_font="B Nazanin", latin_font="Times New Roman", size_pt=10.5, bold=is_header)
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)

    doc.save(FINAL_DOCX)
    if os.path.exists(TEMP_DOCX):
        os.remove(TEMP_DOCX)
    print(f"--> Proposal DOCX created successfully: {FINAL_DOCX}")

if __name__ == "__main__":
    build_proposal_docx()
